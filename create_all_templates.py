from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from xml.sax.saxutils import escape

sdt_id_counter = [100]

def next_id():
    sdt_id_counter[0] += 1
    return sdt_id_counter[0]

def create_template(session_type, output_path):
    """Create a CogScreen tracking note template for a specific session type."""
    global sdt_id_counter
    sdt_id_counter = [100]

    doc = Document()

    # ── Page setup (A4) ──
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Fix zoom
    settings = doc.settings.element
    zoom = settings.find(qn('w:zoom'))
    if zoom is not None:
        zoom.set(qn('w:percent'), '100')

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    # ── Helpers ──
    def insert_pBdr(pPr, border_xml):
        after_elements = ['shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
                          'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN',
                          'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
                          'contextualSpacing', 'mirrorIndents', 'jc', 'rPr']
        for child in pPr:
            tag = child.tag.split('}')[1] if '}' in child.tag else child.tag
            if tag in after_elements:
                child.addprevious(border_xml)
                return
        pPr.append(border_xml)

    def add_separator():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        pPr = p._p.get_or_add_pPr()
        insert_pBdr(pPr, parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            '  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="DDDDDD"/>'
            '</w:pBdr>'
        ))

    def add_title(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"*{title}*")
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.name = "Arial"

    def add_text_sdt(placeholder_text, tag=""):
        sid = next_id()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        tag_xml = f'<w:tag {nsdecls("w")} w:val="{tag}"/>' if tag else ''
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}>'
            f'  <w:sdtPr>'
            f'    <w:id w:val="{sid}"/>'
            f'    {tag_xml}'
            f'    <w:placeholder><w:docPart w:val="DefaultPlaceholder_1082065158"/></w:placeholder>'
            f'    <w:showingPlcHdr/>'
            f'  </w:sdtPr>'
            f'  <w:sdtContent>'
            f'    <w:r>'
            f'      <w:rPr>'
            f'        <w:rStyle w:val="PlaceholderText"/>'
            f'        <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>'
            f'        <w:color w:val="808080"/>'
            f'        <w:sz w:val="22"/><w:szCs w:val="22"/>'
            f'      </w:rPr>'
            f'      <w:t xml:space="preserve">{escape(placeholder_text)}</w:t>'
            f'    </w:r>'
            f'  </w:sdtContent>'
            f'</w:sdt>'
        )
        p._p.append(sdt)

    def add_date_sdt():
        sid = next_id()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}>'
            f'  <w:sdtPr>'
            f'    <w:id w:val="{sid}"/>'
            f'    <w:placeholder><w:docPart w:val="DefaultPlaceholder_1082065158"/></w:placeholder>'
            f'    <w:showingPlcHdr/>'
            f'    <w:date>'
            f'      <w:dateFormat w:val="d MMMM yyyy"/>'
            f'      <w:lid w:val="en-AU"/>'
            f'      <w:storeMappedDataAs w:val="dateTime"/>'
            f'      <w:calendar w:val="gregorian"/>'
            f'    </w:date>'
            f'  </w:sdtPr>'
            f'  <w:sdtContent>'
            f'    <w:r>'
            f'      <w:rPr>'
            f'        <w:rStyle w:val="PlaceholderText"/>'
            f'        <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>'
            f'        <w:color w:val="808080"/>'
            f'        <w:sz w:val="22"/><w:szCs w:val="22"/>'
            f'      </w:rPr>'
            f'      <w:t>Click to select date</w:t>'
            f'    </w:r>'
            f'  </w:sdtContent>'
            f'</w:sdt>'
        )
        p._p.append(sdt)

    def make_list_items(options, placeholder):
        xml = f'<w:listItem {nsdecls("w")} w:displayText="{escape(placeholder)}" w:value=""/>'
        for opt in options:
            xml += f'<w:listItem {nsdecls("w")} w:displayText="{escape(opt)}" w:value="{escape(opt)}"/>'
        return xml

    def add_dropdown_sdt(options, placeholder="Choose an item"):
        sid = next_id()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}>'
            f'  <w:sdtPr>'
            f'    <w:id w:val="{sid}"/>'
            f'    <w:placeholder><w:docPart w:val="DefaultPlaceholder_1082065160"/></w:placeholder>'
            f'    <w:showingPlcHdr/>'
            f'    <w:dropDownList>{make_list_items(options, placeholder)}</w:dropDownList>'
            f'  </w:sdtPr>'
            f'  <w:sdtContent>'
            f'    <w:r><w:rPr>'
            f'      <w:rStyle w:val="PlaceholderText"/>'
            f'      <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>'
            f'      <w:color w:val="808080"/>'
            f'      <w:sz w:val="22"/><w:szCs w:val="22"/>'
            f'    </w:rPr><w:t>{escape(placeholder)}</w:t></w:r>'
            f'  </w:sdtContent>'
            f'</w:sdt>'
        )
        p._p.append(sdt)

    def add_combobox_sdt(options, placeholder="Choose or type your own"):
        sid = next_id()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}>'
            f'  <w:sdtPr>'
            f'    <w:id w:val="{sid}"/>'
            f'    <w:placeholder><w:docPart w:val="DefaultPlaceholder_1082065160"/></w:placeholder>'
            f'    <w:showingPlcHdr/>'
            f'    <w:comboBox>{make_list_items(options, placeholder)}</w:comboBox>'
            f'  </w:sdtPr>'
            f'  <w:sdtContent>'
            f'    <w:r><w:rPr>'
            f'      <w:rStyle w:val="PlaceholderText"/>'
            f'      <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>'
            f'      <w:color w:val="808080"/>'
            f'      <w:sz w:val="22"/><w:szCs w:val="22"/>'
            f'    </w:rPr><w:t>{escape(placeholder)}</w:t></w:r>'
            f'  </w:sdtContent>'
            f'</w:sdt>'
        )
        p._p.append(sdt)

    # Compound helpers
    def add_field(title, placeholder):
        add_title(title)
        add_text_sdt(placeholder)
        add_separator()

    def add_date_field(title):
        add_title(title)
        add_date_sdt()
        add_text_sdt("Time: e.g., 10am \u2013 2pm")
        add_separator()

    def add_dropdown_field(title, options, placeholder="Choose an item"):
        add_title(title)
        add_dropdown_sdt(["N/A \u2013 omit from notes"] + options, placeholder)
        add_separator()

    def add_combobox_field(title, options, placeholder="Choose or type your own"):
        add_title(title)
        add_combobox_sdt(["N/A \u2013 omit from notes"] + options, placeholder)
        add_separator()

    def add_section_header(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"\u2014 {text} \u2014")
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        run.font.name = "Arial"

    # ══════════════════════════════════════
    # BUILD DOCUMENT
    # ══════════════════════════════════════

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(2)
    run = title_para.add_run(f"CogScreen Session Tracking Note \u2014 {session_type}")
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    run.font.name = "Arial"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(8)
    run = sub.add_run("HREC #2022.275")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run.font.name = "Arial"

    # Blue line
    sep_p = doc.add_paragraph()
    sep_p.paragraph_format.space_after = Pt(8)
    pPr = sep_p._p.get_or_add_pPr()
    insert_pBdr(pPr, parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        '  <w:bottom w:val="single" w:sz="12" w:space="1" w:color="2E75B6"/>'
        '</w:pBdr>'
    ))

    # Intro (italic)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    run = intro.add_run(
        "YP consented to participate in the CogScreen (HREC #2022.275) research study "
        "as described below. Their participation was supported by their case manager "
        "who was consulted prior to any research contact. A summary of their research "
        "results (a cognitive summary) will be provided to the case manager to share "
        "with YP as per the protocol and consent procedure. The session is described "
        "in detail below:"
    )
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.name = "Arial"

    # ══════════════════════════════════════
    # COMMON FIELDS
    # ══════════════════════════════════════

    add_date_field("Date and time of session:")

    add_field("Location and attendees:", "e.g., Orygen Parkville \u2013 Brittany and Abigail")

    add_field("Appearance of YP:",
        "Note grooming, hygiene, clothing, and any notable physical features such as weight changes, "
        "psychomotor agitation or slowing, or signs of self-neglect. e.g., tidy, well groomed, dressed "
        "in loose-fitting hoodie and jeans; dishevelled, unkempt hair, clothing stained.")

    add_combobox_field("Eye Contact:", [
        "Appropriate throughout",
        "Limited",
        "Avoided",
        "Improved over session",
        "Brief and intermittent",
        "Varied \u2013 see clinical impressions",
    ])

    add_field("Affect:",
        "Affect is the observable emotional expression you see \u2014 facial expression, vocal tone, "
        "body language \u2014 distinct from mood, which is self-reported. Describe the quality (e.g., "
        "euthymic, flat, anxious, dysphoric, bright), range (restricted, broad, labile), congruence "
        "with content, and reactivity (e.g., non-reactive throughout vs brightened when discussing interests).")

    add_combobox_field("Rapport:", [
        "Easily built",
        "Challenging",
        "Established quickly",
        "Warmed up over time",
        "Difficult to establish",
        "Pre-existing from previous session",
    ])

    add_field("Did anything happen prior to the session that is worth recording?",
        "e.g., lateness, confusion on location, lack of contact, difficulty scheduling, nothing of note")

    # ══════════════════════════════════════
    # CONDITIONAL SECTIONS
    # ══════════════════════════════════════

    if session_type == "Day One":
        add_section_header("Day One")

        add_field("Did the YP have any questions about the consent process?",
            "e.g., asked about data storage, wanted to know who sees results, no questions")

        add_dropdown_field("Was the consent form signed without issue?", [
            "Yes \u2013 signed without hesitation",
            "Yes \u2013 but needed extra time to read through",
            "Yes \u2013 required explanation of specific sections",
            "No \u2013 see notes",
        ])

        add_field("Were there any issues with ACE?",
            "How did YP react to conversation about it? What did they say when asked if CM knows? "
            "Are they okay with CM knowing?")

        add_dropdown_field("Did YP answer yes to any of the first three questions?", [
            "No \u2013 denied all three",
            "Yes \u2013 endorsed item 1",
            "Yes \u2013 endorsed item 2",
            "Yes \u2013 endorsed item 3",
            "Yes \u2013 endorsed multiple items",
        ])

    elif session_type == "Day Two":
        add_section_header("Day Two")

        add_field("How did YP react to Neurocog, anything to note?",
            "e.g., many \u2018don\u2019t knows\u2019, lack of effort, giving up easily, many prompts needed, "
            "found word generation and similarities demanding")

        add_field("How was YP\u2019s mood during neurocog?",
            "e.g., frustrated, enthusiastic, bored, irritated, cycled between positive demeanour and "
            "deflated during harder tasks")

        add_field("Specific considerations/accommodations?",
            "e.g., reading aloud, frequent breaks, snacks, adjusted pacing, session split across two visits")

    elif session_type == "Test Retest":
        add_section_header("Test Retest")

        add_field("How did YP react to Neurocog, anything to note?",
            "e.g., many \u2018don\u2019t knows\u2019, lack of effort, giving up easily, many prompts needed, "
            "found word generation and similarities demanding")

        add_field("How was YP\u2019s mood during neurocog?",
            "e.g., frustrated, enthusiastic, bored, irritated, compared experience to previous session")

        add_field("Specific considerations/accommodations?",
            "e.g., reading aloud, frequent breaks, snacks, adjusted pacing")

        add_field("Any notable differences from previous session?",
            "e.g., more confident with tasks, less anxious, improved engagement, performance appeared consistent")

    # ══════════════════════════════════════
    # ALL SESSIONS (resumed)
    # ══════════════════════════════════════
    add_section_header("All Sessions")

    add_combobox_field("YP\u2019s engagement:", [
        "Remained engaged throughout",
        "Warmed up over time",
        "Disengaged",
        "Enthusiastic",
        "Fluctuated",
        "Initially engaged but fatigued toward end",
    ])

    add_field("Clinical impressions:",
        "e.g., affect presentation, possible negative symptoms, cognitive fatigue, psychomotor "
        "slowing, thought disorganisation")

    add_dropdown_field("Did things go according to protocol?", [
        "Yes",
        "No \u2013 see notes below",
    ])

    add_field("Risk Assessment:",
        "e.g., no risk concerns identified; YP disclosed X \u2013 CM notified; see risk management plan")

    add_combobox_field("How did the session end?", [
        "Positive \u2013 YP left in good spirits",
        "Neutral",
        "YP was upset and needed debriefing",
        "YP left abruptly",
        "Session ended early \u2013 see notes",
    ])

    add_field("Was there any follow-up support?",
        "e.g., drove YP home, contacted case manager, arranged follow-up call, no follow-up needed")

    # ══════════════════════════════════════
    # ADMINISTRATION
    # ══════════════════════════════════════
    add_section_header("Administration")

    add_field("Additional Notes / Actions & Plan:", "Free text for anything not captured above")

    add_dropdown_field("Was payment put through?", [
        "Yes \u2013 processed on the day",
        "No \u2013 pending",
    ])

    add_field("Follow-up actions:",
        "e.g., book Day Two, check in with CM, send cognitive summary")
    add_field("Outstanding items:",
        "e.g., incomplete questionnaires, missing consent form, need to reschedule split session")

    doc.save(output_path)
    print(f"  Created: {output_path}")


# ══════════════════════════════════════
# GENERATE ALL THREE TEMPLATES
# ══════════════════════════════════════
if __name__ == "__main__":
    import os
    out_dir = os.path.dirname(os.path.abspath(__file__))
    print("Building templates...")
    create_template("Day One",      os.path.join(out_dir, "CogScreen_DayOne.docx"))
    create_template("Day Two",      os.path.join(out_dir, "CogScreen_DayTwo.docx"))
    create_template("Test Retest",  os.path.join(out_dir, "CogScreen_TestRetest.docx"))
    print("All done!")
